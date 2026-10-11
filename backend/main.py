from __future__ import annotations

from io import BytesIO
from typing import Any
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

app = FastAPI(title="Excel Structure Converter API", version="0.1.0")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173"], allow_methods=["*"], allow_headers=["*"])
MAX_UPLOAD = 20 * 1024 * 1024


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/workbooks/parse")
async def parse_workbook(file: UploadFile = File(...)) -> dict[str, Any]:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "対応形式は .xlsx です")
    raw = await file.read(MAX_UPLOAD + 1)
    if len(raw) > MAX_UPLOAD:
        raise HTTPException(413, "ファイルサイズは20MB以下にしてください")
    try:
        wb = load_workbook(BytesIO(raw), data_only=False, read_only=False)
        cached = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(400, "Excelファイルを読み込めませんでした") from exc

    sheets = []
    for ws in wb.worksheets:
        cached_ws = cached[ws.title]
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                value = cell.value
                formula = value if isinstance(value, str) and value.startswith("=") else None
                displayed = cached_ws[cell.coordinate].value if formula else value
                fill = cell.fill.fgColor.rgb if cell.fill.fgColor.type == "rgb" else None
                cells.append({"row": cell.row, "column": cell.column, "address": cell.coordinate,
                    "value": displayed if isinstance(displayed, (str, int, float, bool)) else (displayed.isoformat() if displayed else None),
                    "formula": formula, "numberFormat": cell.number_format,
                    "style": {"bold": bool(cell.font.bold), "italic": bool(cell.font.italic), "background": fill,
                        "horizontal": cell.alignment.horizontal, "vertical": cell.alignment.vertical,
                        "wrap": bool(cell.alignment.wrap_text), "borderBottom": cell.border.bottom.style or None}})
        sheets.append({"id": str(uuid4()), "name": ws.title, "rowCount": ws.max_row, "columnCount": ws.max_column,
            "cells": cells, "mergedRanges": [str(r) for r in ws.merged_cells.ranges],
            "rowHeights": {str(i): ws.row_dimensions[i].height for i in range(1, ws.max_row + 1) if ws.row_dimensions[i].height},
            "columnWidths": {get_column_letter(i): ws.column_dimensions[get_column_letter(i)].width for i in range(1, ws.max_column + 1) if ws.column_dimensions[get_column_letter(i)].width},
            "hiddenRows": [i for i in range(1, ws.max_row + 1) if ws.row_dimensions[i].hidden],
            "hiddenColumns": [get_column_letter(i) for i in range(1, ws.max_column + 1) if ws.column_dimensions[get_column_letter(i)].hidden]} )
    return {"id": str(uuid4()), "filename": file.filename, "sheets": sheets}
